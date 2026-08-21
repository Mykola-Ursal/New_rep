#!/usr/bin/env python3
"""Combine the Mexico Staking QA CSV checklists into a single Excel workbook.

Reads the three `Title;How to test;Expected Result` CSV checklists from
`mexico-staking-qa/csv/` and produces `Mexico_Staking_QA_Master.xlsx` with:

- one worksheet per checklist (Smoke Fast, Break It Tonight, Solana Specific),
  each with added `Status` and `Priority` columns,
- a `Summary` worksheet with per-checklist case counts and an overall status
  breakdown once statuses are filled in.

Usage:
    python3 scripts/build_master_workbook.py
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CSV_DIR = ROOT / "csv"
OUTPUT_PATH = ROOT / "Mexico_Staking_QA_Master.xlsx"

CHECKLISTS = [
    ("Smoke Fast", "mexico_smoke_fast.csv"),
    ("Break It Tonight", "mexico_break_it_tonight.csv"),
    ("Solana Specific", "mexico_solana_specific.csv"),
]

STATUS_OPTIONS = ["Not tested", "Passed", "Failed", "Blocked", "Skipped"]
PRIORITY_OPTIONS = ["P0", "P1", "P2", "P3"]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def read_checklist(filename: str) -> list[dict]:
    path = CSV_DIR / filename
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        return list(reader)


def style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    ws.freeze_panes = "A2"


def add_checklist_sheet(wb: Workbook, title: str, rows: list[dict]):
    ws = wb.create_sheet(title=title[:31])
    headers = ["#", "Title", "How to test", "Expected Result", "Status", "Priority"]
    ws.append(headers)

    status_dv = DataValidation(
        type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True
    )
    priority_dv = DataValidation(
        type="list", formula1=f'"{",".join(PRIORITY_OPTIONS)}"', allow_blank=True
    )
    ws.add_data_validation(status_dv)
    ws.add_data_validation(priority_dv)

    for idx, row in enumerate(rows, start=1):
        ws.append(
            [
                idx,
                row.get("Title", ""),
                row.get("How to test", ""),
                row.get("Expected Result", ""),
                "Not tested",
                "",
            ]
        )
        cell_range_status = f"E{idx + 1}"
        cell_range_priority = f"F{idx + 1}"
        status_dv.add(cell_range_status)
        priority_dv.add(cell_range_priority)
        for col in range(1, 7):
            ws.cell(row=idx + 1, column=col).alignment = WRAP

    style_header(ws, len(headers))

    widths = {"A": 5, "B": 40, "C": 60, "D": 55, "E": 14, "F": 10}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    return ws, len(rows)


def add_summary_sheet(wb: Workbook, counts: list[tuple[str, int]]):
    ws = wb.create_sheet(title="Summary", index=0)
    ws.append(["Mexico Staking (MEX) — QA Master Checklist"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Checklist", "Case Count", "Not tested", "Passed", "Failed", "Blocked", "Skipped"])
    style_header(ws, 7)

    row_idx = 4
    for name, count in counts:
        sheet_ref = f"'{name[:31]}'"
        ws.append(
            [
                name,
                count,
                f'=COUNTIF({sheet_ref}!E2:E{count + 1},"Not tested")',
                f'=COUNTIF({sheet_ref}!E2:E{count + 1},"Passed")',
                f'=COUNTIF({sheet_ref}!E2:E{count + 1},"Failed")',
                f'=COUNTIF({sheet_ref}!E2:E{count + 1},"Blocked")',
                f'=COUNTIF({sheet_ref}!E2:E{count + 1},"Skipped")',
            ]
        )
        row_idx += 1

    total_row = row_idx + 1
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for col in range(2, 8):
        col_letter = get_column_letter(col)
        ws.cell(
            row=total_row,
            column=col,
            value=f"=SUM({col_letter}4:{col_letter}{row_idx - 1})",
        ).font = Font(bold=True)

    ws.append([])
    ws.append(["Open findings not part of the original checklists (see bug_reports/):"])
    ws.append(["BUG-001", "Helius RPC API key exposed in plain text query parameter on the front end", "Critical"])
    ws.append(["BUG-002", "Verify rounding behavior on small/frequent claims (dust)", "Medium — pending verification"])

    widths = {"A": 26, "B": 60, "C": 14, "D": 12, "E": 12, "F": 12, "G": 12}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def main():
    wb = Workbook()
    wb.remove(wb.active)

    counts = []
    for title, filename in CHECKLISTS:
        rows = read_checklist(filename)
        _, count = add_checklist_sheet(wb, title, rows)
        counts.append((title, count))

    add_summary_sheet(wb, counts)

    wb.save(OUTPUT_PATH)
    total = sum(c for _, c in counts)
    print(f"Wrote {OUTPUT_PATH} with {total} test cases across {len(counts)} sheets.")


if __name__ == "__main__":
    main()
